"""The .xlsx export: real numbers in real cells, and a file Excel can open."""

from datetime import date
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from stock_monitor import dashboard, excel_export
from stock_monitor.config import NotificationConfig
from stock_monitor.notifier import NotificationDispatcher
from stock_monitor.store import ClosedPosition, Holding, closed_position_store
from stock_monitor.webhook_server import create_webhook_app

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

REPORT = {
    "positions": [
        {"ticker": "AVGO", "quantity": 9.76, "entry_price": 380.0,
         "current_price": 392.23, "market_value": 3828.16, "pnl_pct": 3.22,
         "pnl_value": 119.36, "atr_pct": 5.8, "sector": "Technology",
         "asset_type": "stock", "price_date": None, "price_source": "quote",
         "price_is_stale": False, "purchase_date": "2026-06-20", "holding_days": 45},
        {"ticker": "SEDG", "quantity": 40, "entry_price": 38.2,
         "current_price": 29.44, "market_value": 1177.6, "pnl_pct": -22.93,
         "pnl_value": -350.4, "atr_pct": 7.4, "sector": "Technology",
         "asset_type": "etf", "price_date": "2026-07-31", "price_source": "bar",
         "price_is_stale": True, "purchase_date": None, "holding_days": None},
    ],
    "total_value": 5005.76, "total_pnl_value": -231.04,
    "skipped_tickers": ["ARYT"], "skip_reason": "no price data",
    "volatility": {"exposure_pct": 42.7, "threshold_pct": 30.0},
    "sector": {"threshold_pct": 40.0, "sectors": [
        {"sector": "Technology", "market_value": 5005.76, "weight_pct": 100.0}]},
    "allocation": {"by_asset_type": {"etf_pct": 23.5, "stock_pct": 76.5}},
}

JOURNAL = {
    "entries": [
        {"id": 1, "ticker": "CF", "quantity": 8, "entry_price": 75.0,
         "exit_price": 92.0, "purchase_date": "2026-05-05", "sold_date": "2026-08-03",
         "holding_days": 90, "pnl_pct": 22.67, "pnl_value": 136.0,
         "fraction_sold": 0.4, "is_partial": True, "atr_pct_at_close": 3.1,
         "sector": "Basic Materials", "rating": "green",
         "ai_analysis": "יציאה מתוזמנת היטב", "personal_note": "הערה"},
    ],
    "summary": {"count": 1, "total_pnl": 136.0, "win_rate_pct": 100.0,
                "avg_holding_days": 90},
}


@pytest.fixture
def workbook():
    return load_workbook(BytesIO(excel_export.build_workbook(REPORT, JOURNAL)))


# ── Structure ─────────────────────────────────────────────────────────────────

def test_the_workbook_is_the_holdings_table_and_the_journal(workbook):
    assert workbook.sheetnames == ["התיק שלי", "יומן מסחר"]


def test_the_columns_mirror_the_dashboard_table(workbook):
    """The export exists to be that table, so the layout follows the screen."""
    assert [c.value for c in workbook["התיק שלי"][1]] == [
        "סמל", "כמות", "מחיר כניסה", "מחיר נוכחי", "מטבע", "נכון לתאריך",
        "מקור המחיר", "פרי / פוסט", "שינוי פרי / פוסט (%)", "שווי ($)",
        "רווח/הפסד ($)", "רווח/הפסד (%)", "ימי החזקה", "ATR (%)", "סקטור",
        "תאריך קנייה",
    ]


def test_sheets_are_right_to_left(workbook):
    assert all(workbook[name].sheet_view.rightToLeft for name in workbook.sheetnames)


def test_the_header_row_is_frozen_and_filterable(workbook):
    sheet = workbook["התיק שלי"]
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref is not None


# ── Values ────────────────────────────────────────────────────────────────────

def _row(sheet, ticker):
    headers = [c.value for c in sheet[1]]
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0] == ticker:
            return dict(zip(headers, row))
    raise AssertionError(f"{ticker} not in sheet")


def test_positions_carry_their_numbers(workbook):
    row = _row(workbook["התיק שלי"], "AVGO")
    assert row["כמות"] == pytest.approx(9.76)
    assert row["מחיר נוכחי"] == pytest.approx(392.23)
    assert row["רווח/הפסד ($)"] == pytest.approx(119.36)
    assert row["ימי החזקה"] == 45


def test_numbers_are_numbers_not_text(workbook):
    """The file should be something you can total and pivot, not just read."""
    row = _row(workbook["התיק שלי"], "AVGO")
    for column in ("כמות", "מחיר כניסה", "מחיר נוכחי", "שווי ($)",
                   "רווח/הפסד ($)", "רווח/הפסד (%)", "ATR (%)"):
        assert isinstance(row[column], (int, float)), f"{column} came through as text"


def test_a_loss_stays_negative(workbook):
    assert _row(workbook["התיק שלי"], "SEDG")["רווח/הפסד ($)"] == pytest.approx(-350.4)


def test_dates_are_real_dates(workbook):
    """So Excel can sort and filter them rather than treating them as strings."""
    row = _row(workbook["התיק שלי"], "AVGO")
    assert row["תאריך קנייה"].date() == date(2026, 6, 20)


def test_the_price_source_is_recorded(workbook):
    """A live quote, a close, and an old close must not look identical here.

    The screen distinguishes them with colour and a label; a spreadsheet
    outlives the session, so it needs the distinction in words.
    """
    sheet = workbook["התיק שלי"]
    assert _row(sheet, "AVGO")["מקור המחיר"] == "ציטוט חי"
    assert _row(sheet, "SEDG")["מקור המחיר"] == "סגירה (מסשן קודם)"


def test_a_current_close_is_not_flagged_as_old():
    fresh = {"positions": [dict(REPORT["positions"][1], price_is_stale=False)]}
    book = load_workbook(BytesIO(excel_export.build_workbook(
        fresh, {"entries": [], "summary": {}})))
    assert _row(book["התיק שלי"], "SEDG")["מקור המחיר"] == "סגירה"


def test_the_pre_post_market_columns_are_carried(workbook):
    report = {"positions": [dict(REPORT["positions"][0],
                                 extended_price=394.1, extended_change_pct=0.48)]}
    book = load_workbook(BytesIO(excel_export.build_workbook(
        report, {"entries": [], "summary": {}})))
    row = _row(book["התיק שלי"], "AVGO")
    assert row["פרי / פוסט"] == pytest.approx(394.1)
    assert row["שינוי פרי / פוסט (%)"] == pytest.approx(0.48)


def test_the_journal_sheet_carries_the_ai_verdict(workbook):
    row = _row(workbook["יומן מסחר"], "CF")
    assert row["דירוג"] == "ירוק — החלטה טובה"
    assert row["ניתוח AI"] == "יציאה מתוזמנת היטב"
    assert row["חוות דעת אישית"] == "הערה"
    assert row["מכירה חלקית"] == "40%"


def test_skipped_tickers_are_noted_under_the_table(workbook):
    """Otherwise the totals quietly exclude them and the file looks complete."""
    text = " ".join(
        str(c.value) for row in workbook["התיק שלי"].iter_rows() for c in row
        if c.value is not None
    )
    assert "ARYT" in text


def test_a_totals_row_closes_the_table(workbook):
    """Mirrors the portfolio total shown above the table on screen."""
    sheet = workbook["התיק שלי"]
    headers = [c.value for c in sheet[1]]
    labels = [sheet.cell(row=r, column=1).value for r in range(2, sheet.max_row + 1)]
    total_row = labels.index('סה"כ') + 2

    for column in ("שווי ($)", "רווח/הפסד ($)"):
        cell = sheet.cell(row=total_row, column=headers.index(column) + 1)
        # A live formula, so filtering or editing rows keeps the total honest.
        assert str(cell.value).startswith("=SUM(")


# ── Formatting ────────────────────────────────────────────────────────────────

def test_money_columns_have_a_money_format(workbook):
    sheet = workbook["התיק שלי"]
    headers = [c.value for c in sheet[1]]
    column = headers.index("שווי ($)") + 1
    assert sheet.cell(row=2, column=column).number_format == '#,##0.00'


# ── Edge cases ────────────────────────────────────────────────────────────────

def test_an_empty_portfolio_still_produces_a_valid_file():
    empty = {"positions": [], "total_value": 0, "total_pnl_value": 0}
    book = load_workbook(BytesIO(excel_export.build_workbook(
        empty, {"entries": [], "summary": {}})))
    assert "התיק שלי" in book.sheetnames
    assert book["התיק שלי"].max_row >= 1     # headers survive


def test_missing_optional_fields_do_not_break_the_build():
    sparse = {"positions": [{"ticker": "X", "quantity": 1, "entry_price": 1.0,
                             "current_price": 1.0, "market_value": 1.0,
                             "pnl_pct": 0.0, "pnl_value": 0.0}]}
    assert excel_export.build_workbook(sparse, {"entries": [], "summary": {}})


def test_an_unparsable_date_becomes_blank_rather_than_raising():
    report = {"positions": [dict(REPORT["positions"][0], purchase_date="not a date")]}
    book = load_workbook(BytesIO(excel_export.build_workbook(
        report, {"entries": [], "summary": {}})))
    assert _row(book["התיק שלי"], "AVGO")["תאריך קנייה"] is None


def test_the_filename_carries_the_date():
    assert excel_export.filename(date(2026, 8, 4)) == "portfolio-2026-08-04.xlsx"


# ── Endpoint ──────────────────────────────────────────────────────────────────

@pytest.fixture
def client(monkeypatch):
    monkeypatch.setattr(dashboard._risk_analyzer, "full_report", lambda: REPORT)
    monkeypatch.setattr(dashboard, "get_data_feed", lambda: None)
    return TestClient(create_webhook_app(NotificationDispatcher(NotificationConfig()), ""))


def test_the_endpoint_returns_a_downloadable_workbook(client):
    response = client.get("/api/portfolio/export")

    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX_MIME
    assert "attachment" in response.headers["content-disposition"]
    assert ".xlsx" in response.headers["content-disposition"]
    # It really is a zip-based xlsx, not an error page with the right header.
    assert response.content[:2] == b"PK"
    load_workbook(BytesIO(response.content))


def test_the_export_includes_the_live_journal(client):
    entry = closed_position_store.add(ClosedPosition.from_sale(
        Holding.create("MU", 30, 112.4, purchase_date="2026-07-01"),
        quantity_sold=30, exit_price=98.7))
    try:
        response = client.get("/api/portfolio/export")
        book = load_workbook(BytesIO(response.content))
        tickers = [r[0] for r in book["יומן מסחר"].iter_rows(min_row=2, values_only=True)]
        assert "MU" in tickers
    finally:
        closed_position_store._rows.clear()
        closed_position_store._next_id = 1


def test_a_failure_reports_the_exception_type(client, monkeypatch):
    def boom():
        raise ZeroDivisionError("nope")

    monkeypatch.setattr(dashboard._risk_analyzer, "full_report", boom)
    response = client.get("/api/portfolio/export")
    assert response.status_code == 502
    assert "ZeroDivisionError" in response.json()["detail"]


def test_the_export_is_not_public():
    from stock_monitor.webhook_server import _PUBLIC_PATHS

    assert not any("/api/portfolio/export".startswith(p) for p in _PUBLIC_PATHS)


def test_the_totals_formula_spans_every_data_row():
    """A range that stops short would under-report the portfolio silently."""
    import re

    report = {"positions": [
        dict(REPORT["positions"][0], ticker=t, market_value=100.0, pnl_value=10.0)
        for t in ("AAA", "BBB", "CCC", "DDD")
    ]}
    book = load_workbook(BytesIO(excel_export.build_workbook(
        report, {"entries": [], "summary": {}})))
    sheet = book["התיק שלי"]
    headers = [c.value for c in sheet[1]]
    column = headers.index("שווי ($)") + 1
    labels = [sheet.cell(row=r, column=1).value for r in range(2, sheet.max_row + 1)]
    total_row = labels.index('סה"כ') + 2

    formula = sheet.cell(row=total_row, column=column).value
    first, last = re.search(r"=SUM\([A-Z]+(\d+):[A-Z]+(\d+)\)", formula).groups()
    assert int(first) == 2, "the range must start at the first data row"
    assert int(last) == total_row - 1, "the range must reach the last data row"
    assert int(last) - int(first) + 1 == 4
